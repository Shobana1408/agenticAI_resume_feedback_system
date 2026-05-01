import re
import io
import pdfplumber
import fitz
import pytesseract
from PIL import Image
from docx import Document

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

KNOWN_SKILLS = [
    "python", "java", "c++", "c", "sql", "html", "css", "javascript",
    "flask", "django", "mysql", "git", "github", "react", "node.js",
    "machine learning", "data structures", "problem solving", "rest api",
    "excel", "statistics", "data visualization", "power bi", "tableau",
    "prompt engineering", "llm integration", "api integration",
    "workflow automation", "generative modeling",
    "natural language processing", "machine learning algorithms",
    "data synthesis techniques", "parallel processing",
    "langchain", "langgraph", "rag pipelines", "vector databases",
    "firebase", "figma", "android studio", "flutter", "dart",
    "openai api", "gemini api", "business intelligence",
    "dashboard development", "exploratory data analysis",
    "statistical analysis", "data cleaning",
    "agentic workflows", "tool use", "openai", "gemini"
]

ROLE_WORDS = [
    "software engineer", "data analyst", "agentic ai", "developer",
    "flutter developer", "flutter engineer", "genai", "builder",
    "intern", "engineer", "designer", "analyst", "lead", "member"
]


def format_skill(skill):
    special_map = {
        "sql": "SQL",
        "html": "HTML",
        "css": "CSS",
        "javascript": "JavaScript",
        "c++": "C++",
        "c": "C",
        "mysql": "MySQL",
        "git": "Git",
        "github": "GitHub",
        "react": "React",
        "node.js": "Node.js",
        "rest api": "REST API",
        "power bi": "Power BI",
        "llm integration": "LLM Integration",
        "api integration": "API Integration",
        "natural language processing": "Natural Language Processing",
        "machine learning algorithms": "Machine Learning Algorithms",
        "data synthesis techniques": "Data Synthesis Techniques",
        "parallel processing": "Parallel Processing",
        "generative modeling": "Generative Modeling",
        "workflow automation": "Workflow Automation",
        "prompt engineering": "Prompt Engineering",
        "langchain": "LangChain",
        "langgraph": "LangGraph",
        "rag pipelines": "RAG Pipelines",
        "vector databases": "Vector Databases",
        "firebase": "Firebase",
        "figma": "Figma",
        "android studio": "Android Studio",
        "flutter": "Flutter",
        "dart": "Dart",
        "openai api": "OpenAI API",
        "gemini api": "Gemini API",
        "business intelligence": "Business Intelligence",
        "dashboard development": "Dashboard Development",
        "exploratory data analysis": "Exploratory Data Analysis",
        "statistical analysis": "Statistical Analysis",
        "data cleaning": "Data Cleaning",
        "agentic workflows": "Agentic Workflows",
        "tool use": "Tool Use",
        "openai": "OpenAI",
        "gemini": "Gemini"
    }
    return special_map.get(skill.lower(), skill.title())


def clean_text(text):
    if not text:
        return ""
    text = text.replace("\xa0", " ")
    text = text.replace("\u200b", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n+", "\n", text)
    return text.strip()


def extract_text_from_pdf(file_path):
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text(x_tolerance=2, y_tolerance=2)
                if page_text:
                    text += page_text + "\n"
    except Exception:
        pass
    return clean_text(text)


def extract_text_from_pdf_ocr(file_path):
    text = ""
    try:
        pdf_document = fitz.open(file_path)
        for page_num in range(len(pdf_document)):
            page = pdf_document.load_page(page_num)
            pix = page.get_pixmap(dpi=220)
            img_bytes = pix.tobytes("png")
            image = Image.open(io.BytesIO(img_bytes))
            page_text = pytesseract.image_to_string(image)
            if page_text:
                text += page_text + "\n"
        pdf_document.close()
    except Exception:
        pass
    return clean_text(text)


def extract_text_from_table(table):
    rows_text = []
    for row in table.rows:
        row_items = []
        for cell in row.cells:
            cell_text = clean_text(cell.text)
            if cell_text:
                row_items.append(cell_text)
        if row_items:
            rows_text.append(" | ".join(row_items))
    return "\n".join(rows_text)


def extract_text_from_docx(file_path):
    text_parts = []

    try:
        doc = Document(file_path)

        for para in doc.paragraphs:
            para_text = clean_text(para.text)
            if para_text:
                text_parts.append(para_text)

        for table in doc.tables:
            table_text = extract_text_from_table(table)
            if table_text:
                text_parts.append(table_text)

        for section in doc.sections:
            for para in section.header.paragraphs:
                para_text = clean_text(para.text)
                if para_text:
                    text_parts.append(para_text)

            for para in section.footer.paragraphs:
                para_text = clean_text(para.text)
                if para_text:
                    text_parts.append(para_text)

            for table in section.header.tables:
                table_text = extract_text_from_table(table)
                if table_text:
                    text_parts.append(table_text)

            for table in section.footer.tables:
                table_text = extract_text_from_table(table)
                if table_text:
                    text_parts.append(table_text)

    except Exception:
        pass

    return clean_text("\n".join(text_parts))


def extract_resume_text(file_path):
    if file_path.lower().endswith(".pdf"):
        text = extract_text_from_pdf(file_path)
        if not text or len(text) < 40:
            text = extract_text_from_pdf_ocr(file_path)
        return clean_text(text)

    if file_path.lower().endswith(".docx"):
        return extract_text_from_docx(file_path)

    return ""


def extract_email(text):
    if not text:
        return "Not found"

    match = re.search(r"\b[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}\b", text)
    return match.group(0) if match else "Not found"


def extract_phone(text):
    if not text:
        return "Not found"

    normalized = text.replace("\n", " ")
    normalized = re.sub(r"\s+", " ", normalized)

    patterns = [
        r"(?:\+91[\s\-|]*)?[6-9]\d{9}",
        r"(?:\+91[\s\-|]*)?[6-9]\d{4}[\s\-|]*\d{5}",
        r"\(\+91\)[\s\-|]*[6-9]\d{9}",
        r"\+91[\s\-|]*[6-9]\d{4}[\s\-|]*\d{5}"
    ]

    for pattern in patterns:
        match = re.search(pattern, normalized)
        if match:
            phone = match.group(0).strip()
            phone = phone.replace("|", " ").strip()
            phone = re.sub(r"\s+", " ", phone)
            return phone

    return "Not found"


def is_probable_name_line(line):
    line = clean_text(line)
    if not line:
        return False

    lower_line = line.lower()

    if "@" in line or "|" in line or "http" in lower_line or "www" in lower_line:
        return False

    ignore_words = [
        "resume", "linkedin", "github", "summary", "objective", "education",
        "skills", "experience", "projects", "certifications", "interests",
        "phone", "email", "location", "contact", "references", "clubs",
        "tools", "languages", "about me", "work experience", "ai / genai skills",
        "dev skills", "programming", "software", "projects", "activities"
    ]

    if any(word in lower_line for word in ignore_words):
        return False

    if any(role_word in lower_line for role_word in ROLE_WORDS):
        return False

    if re.search(r"\d", line):
        return False

    words = line.replace("·", " ").replace("•", " ").split()
    if len(words) < 2 or len(words) > 4:
        return False

    clean_words = []
    for word in words:
        word = word.strip(".,-")
        if not re.fullmatch(r"[A-Za-z]+", word):
            return False
        clean_words.append(word)

    return len(clean_words) >= 2


def score_name_candidate(line, index):
    score = 0
    line = clean_text(line)
    words = [w.strip(".,-") for w in line.split()]

    if 2 <= len(words) <= 3:
        score += 3

    if all(w.isupper() for w in words):
        score += 4

    if all(w[:1].isupper() for w in words):
        score += 2

    if index < 5:
        score += 4
    elif index < 10:
        score += 2

    if len(line) <= 30:
        score += 1

    return score


def normalize_name(line):
    words = [w.strip(".,-") for w in clean_text(line).split()]
    normalized_words = []

    for word in words:
        if word.isupper():
            normalized_words.append(word.title())
        else:
            normalized_words.append(word[:1].upper() + word[1:].lower())

    return " ".join(normalized_words)


def extract_name(text):
    if not text:
        return "Not found"

    lines = [clean_text(line) for line in text.splitlines() if clean_text(line)]

    candidates = []

    for index, line in enumerate(lines[:20]):
        if is_probable_name_line(line):
            score = score_name_candidate(line, index)
            candidates.append((score, line))

    if candidates:
        candidates.sort(key=lambda x: x[0], reverse=True)
        return normalize_name(candidates[0][1])

    return "Not found"


def extract_skills(text):
    if not text:
        return []

    text_lower = text.lower()
    found_skills = []

    sorted_skills = sorted(KNOWN_SKILLS, key=len, reverse=True)

    for skill in sorted_skills:
        pattern = r"\b" + re.escape(skill.lower()) + r"\b"
        if re.search(pattern, text_lower):
            formatted = format_skill(skill)
            if formatted not in found_skills:
                found_skills.append(formatted)

    return found_skills


def parse_resume(file_path):
    resume_text = extract_resume_text(file_path)

    parsed_data = {
        "name": extract_name(resume_text),
        "email": extract_email(resume_text),
        "phone": extract_phone(resume_text),
        "skills": extract_skills(resume_text),
        "resume_text": resume_text
    }

    return parsed_data